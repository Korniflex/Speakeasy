import tkinter as tk
from tkinter import filedialog, messagebox
from PIL import Image, ImageTk
import pandas as pd
import PyPDF2
import pytesseract
from docx import Document
import os

class SpeakeasyApp:
    """
    Speakeasy Text Extractor Application
    
    A GUI application that extracts text from various file formats including:
    - Excel (.xlsx)
    - Word (.docx)
    - Plain text (.txt)
    - CSV (.csv)
    - PDF (.pdf)
    - Images (.jpg, .jpeg, .png, .gif)
    """
    
    def __init__(self, root):
        """Initialize the application."""
        self.root = root
        self.root.geometry("550x400")
        self.root.title("Speakeasy v0.1")
        self.root.resizable(width=True, height=True)
        
        self.setup_ui()
        
    def setup_ui(self):
        """Set up the user interface."""
        # Load and display logo
        try:
            logo_path = os.path.join(os.path.dirname(__file__), "data", "logo_Speakeasy_1080p.png")
            logo = Image.open(logo_path)
            logo = logo.resize((200, 200))
            self.logo_image = ImageTk.PhotoImage(logo)
            
            # Create logo label
            self.logo_label = tk.Label(
                self.root,
                text='Willkommen bei Speakeasy.\nBitte wählen Sie eine Datei aus, um den Text zu extrahieren',
                image=self.logo_image,
                compound='top'
            )
            self.logo_label.pack(anchor=tk.CENTER, pady=10)
        except Exception as e:
            print(f"Fehler beim Laden des Logos: {e}")
            # Create text-only label if logo fails to load
            self.logo_label = tk.Label(
                self.root,
                text='Willkommen bei Speakeasy.\nBitte wählen Sie eine Datei aus, um den Text zu extrahieren',
                font=('Arial', 14)
            )
            self.logo_label.pack(anchor=tk.CENTER, pady=20)
        
        # Create buttons
        button_frame = tk.Frame(self.root)
        button_frame.pack(pady=20)
        
        self.open_button = tk.Button(
            button_frame,
            text="Datei öffnen",
            command=self.process_file,
            width=15,
            height=2
        )
        self.open_button.grid(row=0, column=0, padx=10)
        
        self.save_button = tk.Button(
            button_frame,
            text="Als Text speichern",
            command=lambda: self.save_to_text(self.extracted_text),
            width=15,
            height=2,
            state=tk.DISABLED
        )
        self.save_button.grid(row=0, column=1, padx=10)
        
        # Text display area
        self.text_frame = tk.Frame(self.root)
        self.text_frame.pack(fill=tk.BOTH, expand=True, padx=20, pady=10)
        
        self.text_display = tk.Text(self.text_frame, wrap=tk.WORD, height=10)
        self.text_display.pack(fill=tk.BOTH, expand=True)
        
        # Scrollbar for text display
        scrollbar = tk.Scrollbar(self.text_display)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        self.text_display.config(yscrollcommand=scrollbar.set)
        scrollbar.config(command=self.text_display.yview)
        
        # Initialize extracted text variable
        self.extracted_text = None
        
    def process_file(self):
        """Open a file and extract its text."""
        file_path = self.open_file()
        if file_path:
            self.extracted_text = self.read_file(file_path)
            if self.extracted_text:
                # Display the extracted text
                self.text_display.delete(1.0, tk.END)
                self.text_display.insert(tk.END, self.extracted_text[:1000] + 
                                        ("\n\n[...] Text gekürzt für die Anzeige" if len(self.extracted_text) > 1000 else ""))
                # Enable save button
                self.save_button.config(state=tk.NORMAL)
            else:
                messagebox.showerror("Fehler", "Konnte keinen Text aus der Datei extrahieren.")
    
    def open_file(self):
        """Open a file dialog and return the selected file path."""
        file_types = [
            ("Alle unterstützten Dateien", "*.xlsx;*.docx;*.txt;*.csv;*.pdf;*.jpg;*.jpeg;*.png;*.gif"),
            ("Excel-Dateien", "*.xlsx"),
            ("Word-Dokumente", "*.docx"),
            ("Text-Dateien", "*.txt"),
            ("CSV-Dateien", "*.csv"),
            ("PDF-Dateien", "*.pdf"),
            ("Bild-Dateien", "*.jpg;*.jpeg;*.png;*.gif")
        ]
        file_path = filedialog.askopenfilename(filetypes=file_types)
        return file_path
    
    def read_file(self, file_path):
        """Read and extract text from a file based on its extension."""
        if not file_path:
            return None
        
        file_extension = os.path.splitext(file_path)[1].lower()
        
        readers = {
            '.xlsx': self.read_excel,
            '.docx': self.read_word,
            '.txt': self.read_txt,
            '.csv': self.read_csv,
            '.pdf': self.read_pdf,
            '.jpg': self.read_image,
            '.jpeg': self.read_image,
            '.png': self.read_image,
            '.gif': self.read_image
        }
        
        if file_extension in readers:
            try:
                return readers[file_extension](file_path)
            except Exception as e:
                messagebox.showerror("Fehler", f"Fehler beim Lesen der Datei: {e}")
                return None
        else:
            messagebox.showerror("Fehler", "Dateiformat wird nicht unterstützt.")
            return None
    
    def read_excel(self, file_path):
        """Extract text from an Excel file."""
        try:
            df = pd.read_excel(file_path)
            return df.to_string()
        except Exception as e:
            print(f'Fehler beim Lesen der Excel-Datei: {e}')
            return None
    
    def read_word(self, file_path):
        """Extract text from a Word document."""
        try:
            document = Document(file_path)
            text = []
            for paragraph in document.paragraphs:
                text.append(paragraph.text)
            return "\n".join(text)
        except Exception as e:
            print(f'Fehler beim Lesen des Word-Dokuments: {e}')
            return None
    
    def read_csv(self, file_path):
        """Extract text from a CSV file."""
        try:
            df = pd.read_csv(file_path)
            return df.to_string()
        except Exception as e:
            print(f'Fehler beim Lesen der CSV-Datei: {e}')
            return None
    
    def read_txt(self, file_path):
        """Extract text from a plain text file."""
        try:
            with open(file_path, 'r', encoding='utf-8') as file:
                text = file.read()
            return text
        except UnicodeDecodeError:
            # Try with a different encoding if UTF-8 fails
            try:
                with open(file_path, 'r', encoding='latin-1') as file:
                    text = file.read()
                return text
            except Exception as e:
                print(f'Fehler beim Lesen der Textdatei: {e}')
                return None
        except Exception as e:
            print(f'Fehler beim Lesen der Textdatei: {e}')
            return None
    
    def read_pdf(self, file_path):
        """Extract text from a PDF file."""
        try:
            with open(file_path, "rb") as file:
                reader = PyPDF2.PdfReader(file)
                text = ""
                for page_num in range(len(reader.pages)):
                    page = reader.pages[page_num]
                    text += page.extract_text() or ""  # Handle empty pages
                return text
        except Exception as e:
            print(f"Fehler beim Extrahieren des Textes aus der PDF: {e}")
            return None
    
    def read_image(self, file_path):
        """Extract text from an image using OCR."""
        try:
            img = Image.open(file_path)
            text = pytesseract.image_to_string(img)
            return text
        except Exception as e:
            print(f"Fehler beim Extrahieren des Textes aus dem Bild: {e}")
            return None
    
    def save_to_text(self, text):
        """Save the extracted text to a file."""
        if not text:
            messagebox.showerror("Fehler", "Kein Text zum Speichern verfügbar.")
            return
        
        output_file = filedialog.asksaveasfilename(
            defaultextension=".txt",
            filetypes=[("Text-Dateien", "*.txt")]
        )
        
        if output_file:
            try:
                with open(output_file, "w", encoding="utf-8") as f:
                    f.write(text)
                messagebox.showinfo("Erfolg", f"Text gespeichert in {output_file}")
            except Exception as e:
                messagebox.showerror("Fehler", f"Fehler bei der Speicherung: {e}")


if __name__ == "__main__":
    root = tk.Tk()
    app = SpeakeasyApp(root)
    root.mainloop()